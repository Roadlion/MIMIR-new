# backend/app/utils/docx_skill.py
"""
MIMIR Agentic Create DOCX Skill Engine
Inspired by Claude / Anthropic professional document generation skills.

Converts structured markdown, research session history, and financial memos into
institutional-grade Word (.docx) documents with custom typography, branded colors,
shaded tables, callout boxes, and headers/footers.
"""

import io
import re
from datetime import datetime
from typing import List, Dict, Any, Optional

import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

# --- MIMIR Brand Color Palette ---
COLOR_TEAL_DARK = RGBColor(0, 102, 112)     # #006670
COLOR_TEAL_ACCENT = RGBColor(0, 166, 178)  # #00A6B2
COLOR_CHARCOAL = RGBColor(26, 42, 48)      # #1A2A30
COLOR_DARK_TEXT = RGBColor(12, 20, 23)     # #0C1417
COLOR_MUTED = RGBColor(74, 106, 112)       # #4A6A70
COLOR_WHITE = RGBColor(255, 255, 255)

HEX_TEAL_DARK = "006670"
HEX_TEAL_ACCENT = "00A6B2"
HEX_ZEBRA_LIGHT = "F4F7F7"
HEX_BORDER_LIGHT = "D0DCDD"
HEX_CALLOUT_BG = "EEF6F6"
HEX_CALLOUT_BORDER = "00A6B2"
HEX_CODE_BG = "F0F4F4"

# --- OXML Formatting Utilities ---
def set_cell_background(cell, fill_hex: str):
    """Set background shading color of a table cell."""
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=180, right=180):
    """Set inner cell padding (in dxa: 20 dxa = 1 pt)."""
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_cell_borders(cell, top=None, bottom=None, left=None, right=None):
    """Set borders for a specific table cell."""
    tcPr = cell._element.get_or_add_tcPr()
    tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}/>')
    
    borders = {'top': top, 'bottom': bottom, 'left': left, 'right': right}
    for side, b_data in borders.items():
        if b_data:
            val = b_data.get("val", "single")
            sz = b_data.get("sz", 4)
            color = b_data.get("color", "auto")
            b_xml = parse_xml(f'<w:{side} {nsdecls("w")} w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>')
            tcBorders.append(b_xml)
        else:
            b_xml = parse_xml(f'<w:{side} {nsdecls("w")} w:val="none"/>')
            tcBorders.append(b_xml)
    tcPr.append(tcBorders)

def add_callout_box(doc: Document, title: str, text: str, alert_type: str = "NOTE"):
    """Render a styled callout / alert box with left accent border and shaded fill."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    
    set_cell_background(cell, HEX_CALLOUT_BG)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    color_hex = HEX_TEAL_ACCENT
    if alert_type in ("IMPORTANT", "WARNING", "CAUTION"):
        color_hex = "FF5252" if alert_type == "CAUTION" else "FF9800"
        
    set_cell_borders(cell, left={"val": "single", "sz": 24, "color": color_hex},
                           top={"val": "none"}, bottom={"val": "none"}, right={"val": "none"})
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    run_t = p.add_run(f"[{alert_type}] {title}\n" if title else f"[{alert_type}]\n")
    run_t.bold = True
    run_t.font.size = Pt(10)
    run_t.font.color.rgb = COLOR_TEAL_DARK
    
    process_inline_formatting(p, text)
    doc.add_paragraph()  # spacing after table

def process_inline_formatting(paragraph, text: str):
    """Parses **bold**, *italic*, and `code` inline formatting."""
    tokens = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`)', text)
    for token in tokens:
        if not token:
            continue
        if token.startswith('**') and token.endswith('**'):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith('*') and token.endswith('*'):
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        elif token.startswith('`') and token.endswith('`'):
            run = paragraph.add_run(token[1:-1])
            run.font.name = 'Consolas'
            run.font.size = Pt(9.5)
            run.font.color.rgb = COLOR_TEAL_DARK
        else:
            paragraph.add_run(token)

def render_markdown_table(doc: Document, table_lines: List[str]):
    """Parses markdown table lines (| Col 1 | Col 2 |) into a styled docx Table."""
    rows_data = []
    for line in table_lines:
        line_strip = line.strip().strip('|')
        if not line_strip:
            continue
        # Skip alignment delimiter row (|---|---|)
        if re.match(r'^[:\-\s|]+$', line_strip):
            continue
        cells = [c.strip() for c in line.strip().split('|')[1:-1]]
        if cells:
            rows_data.append(cells)
            
    if not rows_data:
        return
        
    num_rows = len(rows_data)
    num_cols = max(len(r) for r in rows_data)
    
    table = doc.add_table(rows=num_rows, cols=num_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    for r_idx, row_cells in enumerate(rows_data):
        for c_idx, cell_text in enumerate(row_cells):
            if c_idx < num_cols:
                cell = table.cell(r_idx, c_idx)
                cell.text = ""
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(3)
                p.paragraph_format.space_after = Pt(3)
                
                # Header row styling
                if r_idx == 0:
                    set_cell_background(cell, HEX_TEAL_DARK)
                    set_cell_margins(cell, top=140, bottom=140, left=160, right=160)
                    process_inline_formatting(p, cell_text)
                    for r in p.runs:
                        r.bold = True
                        r.font.color.rgb = COLOR_WHITE
                        r.font.size = Pt(9.5)
                else:
                    # Zebra striping
                    bg = HEX_ZEBRA_LIGHT if (r_idx % 2 == 1) else "FFFFFF"
                    set_cell_background(cell, bg)
                    set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
                    set_cell_borders(cell, 
                                     top={"val": "single", "sz": 2, "color": HEX_BORDER_LIGHT},
                                     bottom={"val": "single", "sz": 2, "color": HEX_BORDER_LIGHT},
                                     left={"val": "none"}, right={"val": "none"})
                    process_inline_formatting(p, cell_text)
                    for r in p.runs:
                        r.font.size = Pt(9.5)
                        r.font.color.rgb = COLOR_DARK_TEXT
                        
    doc.add_paragraph()  # Spacing after table

def build_docx_document(title: str, markdown_content: str, session_meta: Optional[Dict[str, Any]] = None) -> io.BytesIO:
    """
    Main Create DOCX Skill engine entrypoint.
    Converts markdown text into a styled, professional Word document.
    """
    doc = Document()
    
    # 1. Page Margins & Setup
    sections = doc.sections
    for s in sections:
        s.top_margin = Inches(0.9)
        s.bottom_margin = Inches(0.9)
        s.left_margin = Inches(0.9)
        s.right_margin = Inches(0.9)
        
        # Header & Footer setup
        header = s.header
        hp = header.paragraphs[0]
        hp.alignment = WD_PARAGRAPH_ALIGNMENT.RIGHT
        hrun = hp.add_run("MIMIR Oracle Financial Intelligence Report")
        hrun.font.name = "Calibri"
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = COLOR_MUTED
        
        footer = s.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
        frun = fp.add_run(f"Confidential · Generated on {datetime.now().strftime('%B %d, %Y')} by MIMIR Oracle Assistant")
        frun.font.name = "Calibri"
        frun.font.size = Pt(8.5)
        frun.font.color.rgb = COLOR_MUTED
        
    # 2. Cover / Title Block Header
    cover_table = doc.add_table(rows=1, cols=1)
    cover_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_cell = cover_table.cell(0, 0)
    set_cell_background(c_cell, HEX_TEAL_DARK)
    set_cell_margins(c_cell, top=200, bottom=200, left=240, right=240)
    
    cp = c_cell.paragraphs[0]
    cp.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
    run_t = cp.add_run(title.upper() + "\n")
    run_t.bold = True
    run_t.font.name = "Arial"
    run_t.font.size = Pt(20)
    run_t.font.color.rgb = COLOR_WHITE
    
    sub_text = f"INSTITUTIONAL RESEARCH DOSSIER · {datetime.now().strftime('%Y-%m-%d %H:%M')}"
    if session_meta and session_meta.get("session_id"):
        sub_text += f" · SESSION: {session_meta['session_id'][:8]}"
    run_sub = cp.add_run(sub_text)
    run_sub.font.name = "Calibri"
    run_sub.font.size = Pt(9.5)
    run_sub.font.color.rgb = COLOR_TEAL_ACCENT
    
    doc.add_paragraph()  # spacing
    
    # 3. Parse Markdown Line by Line
    lines = markdown_content.split('\n')
    idx = 0
    in_code_block = False
    code_lines = []
    
    while idx < len(lines):
        line = lines[idx]
        line_strip = line.strip()
        
        # Code Block Handling
        if line_strip.startswith('```'):
            if in_code_block:
                # End code block
                if code_lines:
                    p = doc.add_paragraph()
                    p.paragraph_format.space_before = Pt(4)
                    p.paragraph_format.space_after = Pt(6)
                    run = p.add_run('\n'.join(code_lines))
                    run.font.name = 'Consolas'
                    run.font.size = Pt(9)
                    run.font.color.rgb = COLOR_CHARCOAL
                code_lines = []
                in_code_block = False
            else:
                in_code_block = True
                code_lines = []
            idx += 1
            continue
            
        if in_code_block:
            code_lines.append(line)
            idx += 1
            continue
            
        if not line_strip:
            idx += 1
            continue
            
        # Table Handling
        if line_strip.startswith('|') and line_strip.endswith('|'):
            table_lines = []
            while idx < len(lines) and lines[idx].strip().startswith('|') and lines[idx].strip().endswith('|'):
                table_lines.append(lines[idx])
                idx += 1
            render_markdown_table(doc, table_lines)
            continue
            
        # Headings
        if line_strip.startswith('# '):
            h = doc.add_heading(level=1)
            h.paragraph_format.space_before = Pt(16)
            h.paragraph_format.space_after = Pt(6)
            run = h.add_run(line_strip[2:])
            run.bold = True
            run.font.name = "Arial"
            run.font.size = Pt(16)
            run.font.color.rgb = COLOR_TEAL_DARK
        elif line_strip.startswith('## '):
            h = doc.add_heading(level=2)
            h.paragraph_format.space_before = Pt(12)
            h.paragraph_format.space_after = Pt(4)
            run = h.add_run(line_strip[3:])
            run.bold = True
            run.font.name = "Arial"
            run.font.size = Pt(13)
            run.font.color.rgb = COLOR_CHARCOAL
        elif line_strip.startswith('### '):
            h = doc.add_heading(level=3)
            h.paragraph_format.space_before = Pt(10)
            h.paragraph_format.space_after = Pt(4)
            run = h.add_run(line_strip[4:])
            run.bold = True
            run.font.name = "Calibri"
            run.font.size = Pt(11.5)
            run.font.color.rgb = COLOR_TEAL_ACCENT
        # Callouts / Blockquotes
        elif line_strip.startswith('> '):
            callout_text = line_strip[2:]
            alert_type = "NOTE"
            if "[!IMPORTANT]" in callout_text:
                alert_type = "IMPORTANT"
                callout_text = callout_text.replace("[!IMPORTANT]", "")
            elif "[!WARNING]" in callout_text:
                alert_type = "WARNING"
                callout_text = callout_text.replace("[!WARNING]", "")
            add_callout_box(doc, "", callout_text.strip(), alert_type=alert_type)
        # Lists
        elif line_strip.startswith('- ') or line_strip.startswith('* '):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            process_inline_formatting(p, line_strip[2:])
        elif re.match(r'^\d+\.\s', line_strip):
            clean_item = re.sub(r'^\d+\.\s', '', line_strip)
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            process_inline_formatting(p, clean_item)
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(4)
            process_inline_formatting(p, line_strip)
            
        idx += 1
        
    # Save to BytesIO Buffer
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer
